"""Tests for fulltext_eligibility.export (XLSX review workbook)."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook
from pydantic import BaseModel

from laglitsynth.catalogue_fetch.models import Work
from laglitsynth.fulltext_eligibility.export import (
    _unique_sheet_name,
    export_review_xlsx,
    sample_verdicts,
    short_work_id,
)
from laglitsynth.fulltext_eligibility.models import EligibilityVerdict

from conftest import _make_work


def _write_jsonl(path: Path, records: list[BaseModel]) -> None:
    with open(path, "w") as f:
        for r in records:
            f.write(r.model_dump_json() + "\n")


def _write_inputs(
    tmp_path: Path,
    works: list[Work],
    verdicts: list[EligibilityVerdict],
) -> tuple[Path, Path, Path]:
    catalogue_path = tmp_path / "dedup.jsonl"
    verdicts_path = tmp_path / "verdicts.jsonl"
    output_path = tmp_path / "review.xlsx"
    _write_jsonl(catalogue_path, list(works))
    _write_jsonl(verdicts_path, list(verdicts))
    return verdicts_path, catalogue_path, output_path


def _write_meta(tmp_path: Path, criterion: str) -> Path:
    meta_path = tmp_path / "eligibility-meta.json"
    meta_path.write_text(
        '{"run": {"tool": "t", "tool_version": "alpha",'
        ' "run_at": "2026-01-01T00:00:00.000000+00:00", "validation_skipped": 0},'
        ' "llm": {"model": "gemma3:4b", "temperature": 0.8, "prompt_sha256": "'
        + ("a" * 64)
        + '"},'
        ' "input_catalogue": "c", "input_screening_verdicts": "s",'
        ' "input_extractions": "e", "input_count": 1, "eligible_count": 1,'
        ' "excluded_count": 0, "tei_parse_failure_count": 0,'
        ' "llm_parse_failure_count": 0, "llm_timeout_count": 0,'
        ' "criterion": ' + _json_str(criterion) + "}"
    )
    return meta_path


def _json_str(s: str) -> str:
    import json

    return json.dumps(s)


# ── sampling helpers ────────────────────────────────────────────────────────


def test_sample_verdicts_none_returns_all() -> None:
    vs = [EligibilityVerdict(work_id=f"W{i}") for i in range(5)]
    assert sample_verdicts(vs, None, seed=0) == vs


def test_sample_verdicts_preserves_original_order() -> None:
    vs = [EligibilityVerdict(work_id=f"W{i}") for i in range(10)]
    sample = sample_verdicts(vs, 5, seed=42)
    assert len(sample) == 5
    positions = [vs.index(v) for v in sample]
    assert positions == sorted(positions)


def test_sample_verdicts_reproducible() -> None:
    vs = [EligibilityVerdict(work_id=f"W{i}") for i in range(20)]
    a = [v.work_id for v in sample_verdicts(vs, 7, seed=123)]
    b = [v.work_id for v in sample_verdicts(vs, 7, seed=123)]
    c = [v.work_id for v in sample_verdicts(vs, 7, seed=456)]
    assert a == b
    assert a != c


def test_short_work_id_openalex_url() -> None:
    assert short_work_id("https://openalex.org/W3213722062") == "W3213722062"


def test_unique_sheet_name_collision_suffix() -> None:
    used: set[str] = set()
    assert _unique_sheet_name("W1", used) == "W1"
    assert _unique_sheet_name("W1", used) == "W1_2"


# ── overview / Index sheet ──────────────────────────────────────────────────


def test_index_sheet_inline_verdict_and_hyperlink(tmp_path: Path) -> None:
    """Overview sheet has one row per sampled work with eligible/llm_reason
    inline and a hyperlink to the per-work tab."""
    works = [
        _make_work("https://openalex.org/W1"),
        _make_work("https://openalex.org/W2"),
    ]
    verdicts = [
        EligibilityVerdict(
            work_id="https://openalex.org/W1",
            eligible=True,
            reason="meets criterion",
            seed=1,
        ),
        EligibilityVerdict(
            work_id="https://openalex.org/W2",
            eligible=False,
            reason="review article",
            seed=2,
        ),
    ]
    verdicts_path, catalogue_path, output_path = _write_inputs(tmp_path, works, verdicts)

    count = export_review_xlsx(
        verdicts_path, catalogue_path, output_path, n_subset=None
    )

    assert count == 2
    wb = load_workbook(output_path)
    assert wb.sheetnames == ["Index", "W1", "W2"]
    index = wb["Index"]
    # Header row (row 5) carries eligible / llm_reason inline plus sheet link.
    assert index["A5"].value == "work_id"
    assert index["F5"].value == "eligible"
    assert index["G5"].value == "llm_reason"
    assert index["H5"].value == "sheet"
    # W1 eligible=yes, W2 eligible=no.
    assert index["A6"].value == "https://openalex.org/W1"
    assert index["F6"].value == "yes"
    assert index["G6"].value == "meets criterion"
    assert index["F7"].value == "no"
    assert index["G7"].value == "review article"
    # Sheet hyperlink points at the per-work tab.
    link = index["H6"]
    assert link.value == "W1"
    assert link.hyperlink is not None
    assert link.hyperlink.location == "'W1'!A1"


# ── per-work sheet ──────────────────────────────────────────────────────────


def test_work_sheet_embeds_criterion_and_reviewer_placeholders(
    tmp_path: Path,
) -> None:
    work = _make_work(
        "https://openalex.org/W1",
        title="The Title",
        doi="10.1000/xyz",
        publication_year=2024,
    )
    verdicts = [
        EligibilityVerdict(
            work_id="https://openalex.org/W1",
            eligible=True,
            reason="meets criterion",
            seed=1,
            raw_response='{"eligible": true, "reason": "meets criterion"}',
        )
    ]
    verdicts_path, catalogue_path, output_path = _write_inputs(tmp_path, works=[work], verdicts=verdicts)
    meta_path = _write_meta(tmp_path, "Is this primary research on Lagrangian methods?")

    export_review_xlsx(
        verdicts_path,
        catalogue_path,
        output_path,
        n_subset=None,
        meta_path=meta_path,
    )
    wb = load_workbook(output_path)
    ws = wb["W1"]

    # Back-to-Index hyperlink.
    assert ws["A1"].value == "← back to Index"
    assert ws["A1"].hyperlink.location == "'Index'!A1"

    # Bibliographic block.
    assert ws["A3"].value == "title"
    assert ws["B3"].value == "The Title"
    assert ws["A7"].value == "doi"
    assert ws["B7"].hyperlink.target == "https://doi.org/10.1000/xyz"
    assert ws["A8"].value == "openalex"
    assert ws["B8"].hyperlink.target == "https://openalex.org/W1"

    # Criterion embedded verbatim from meta.
    assert ws["A10"].value == "criterion"
    assert ws["B10"].value == "Is this primary research on Lagrangian methods?"

    # Reviewer placeholders.
    assert ws["A11"].value == "reviewer_eligible"
    assert "<enter yes/no" in ws["B11"].value
    assert ws["A12"].value == "reviewer_reason"
    assert ws["A13"].value == "reviewer_remarks"

    # Collapsed LLM block.
    assert ws["A15"].value == "LLM details (expand to peek)"
    assert ws["A16"].value == "llm_eligible"
    assert ws["B16"].value == "yes"
    assert ws["A17"].value == "llm_reason"
    assert ws["B17"].value == "meets criterion"
    for row_idx in range(16, 22):
        assert ws.row_dimensions[row_idx].outline_level == 1
        assert ws.row_dimensions[row_idx].hidden is True
    assert ws.sheet_properties.outlinePr.summaryBelow is False


# ── subset / seed ───────────────────────────────────────────────────────────


def test_subset_samples_reproducibly(tmp_path: Path) -> None:
    works = [_make_work(f"https://openalex.org/W{i}") for i in range(10)]
    verdicts = [
        EligibilityVerdict(
            work_id=f"https://openalex.org/W{i}",
            eligible=True,
            reason=f"r{i}",
            seed=i,
        )
        for i in range(10)
    ]
    verdicts_path, catalogue_path, output_path = _write_inputs(tmp_path, works, verdicts)

    count = export_review_xlsx(
        verdicts_path, catalogue_path, output_path, n_subset=4, seed=7
    )
    assert count == 4
    wb = load_workbook(output_path)
    per_work = wb.sheetnames[1:]
    assert len(per_work) == 4
    # Sheet order follows verdict (ascending) order.
    indices = [int(name[1:]) for name in per_work]
    assert indices == sorted(indices)

    # Re-running with the same seed selects the same works.
    output_path2 = tmp_path / "review2.xlsx"
    export_review_xlsx(
        verdicts_path, catalogue_path, output_path2, n_subset=4, seed=7
    )
    wb2 = load_workbook(output_path2)
    assert wb2.sheetnames[1:] == per_work


# ── missing catalogue ───────────────────────────────────────────────────────


def test_raises_on_missing_work(tmp_path: Path) -> None:
    works = [_make_work("https://openalex.org/W1")]
    verdicts = [
        EligibilityVerdict(
            work_id="https://openalex.org/W_missing",
            eligible=True,
            reason="?",
        )
    ]
    verdicts_path, catalogue_path, output_path = _write_inputs(tmp_path, works, verdicts)

    with pytest.raises(ValueError, match="W_missing"):
        export_review_xlsx(
            verdicts_path, catalogue_path, output_path, n_subset=None
        )


# ── CLI smoke ───────────────────────────────────────────────────────────────


def test_cli_default_output(tmp_path: Path) -> None:
    from laglitsynth.cli import main

    works = [_make_work("https://openalex.org/W1")]
    verdicts = [
        EligibilityVerdict(
            work_id="https://openalex.org/W1",
            eligible=True,
            reason="ok",
            seed=1,
        )
    ]
    verdicts_path, catalogue_path, _ = _write_inputs(tmp_path, works, verdicts)

    main(
        [
            "fulltext-eligibility-export",
            "--verdicts",
            str(verdicts_path),
            "--catalogue",
            str(catalogue_path),
        ]
    )

    assert (tmp_path / "review.xlsx").exists()

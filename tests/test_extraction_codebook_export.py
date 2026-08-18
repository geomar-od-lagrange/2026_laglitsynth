"""Tests for extraction_codebook.export (XLSX review workbook)."""

from __future__ import annotations

from pathlib import Path
from typing import Any, cast

import pytest
from openpyxl import load_workbook
from pydantic import BaseModel

from laglitsynth.catalogue_fetch.models import Work
from laglitsynth.extraction_codebook.codebook import (
    CodebookContext,
    ExtractionRecordProto,
)
from laglitsynth.extraction_codebook.export import (
    _unique_sheet_name,
    export_review_xlsx,
    sample_records,
    short_work_id,
)

from conftest import _make_work

DEFAULT_CODEBOOK_PATH = Path("examples/codebooks/lagrangian-oceanography.yaml")

_ALT_CODEBOOK = {
    "id": "alt",
    "system_prompt": "Extract these fields:\n{fields}",
    "fields": [
        {"name": "alpha", "description": "The alpha field."},
        {"name": "beta", "description": "The beta field.", "context": False},
    ],
}


@pytest.fixture(scope="module")
def ctx() -> CodebookContext:
    return CodebookContext.from_spec(DEFAULT_CODEBOOK_PATH)


def _write_jsonl(path: Path, records: list[BaseModel]) -> None:
    with open(path, "w") as f:
        for r in records:
            f.write(r.model_dump_json() + "\n")


def cast_record(r: BaseModel) -> ExtractionRecordProto:
    return cast(ExtractionRecordProto, r)


def _make_record(ctx: CodebookContext, work_id: str, **values: Any) -> BaseModel:
    payload_fields = {name: None for name in ctx.payload_field_names}
    payload_fields.update(values)
    return ctx.record_model(
        work_id=work_id,
        reason=None,
        seed=1,
        truncated=False,
        raw_response='{"x": "y"}',
        **payload_fields,
    )


def _write_inputs(
    tmp_path: Path,
    works: list[Work],
    records: list[BaseModel],
) -> tuple[Path, Path, Path]:
    catalogue_path = tmp_path / "dedup.jsonl"
    records_path = tmp_path / "records.jsonl"
    output_path = tmp_path / "review.xlsx"
    _write_jsonl(catalogue_path, list(works))
    _write_jsonl(records_path, records)
    return records_path, catalogue_path, output_path


# ── sampling helpers ────────────────────────────────────────────────────────


def test_sample_records_none_returns_all(ctx: CodebookContext) -> None:
    rs = [
        cast_record(_make_record(ctx, f"W{i}")) for i in range(5)
    ]
    assert sample_records(rs, None, seed=0) == rs


def test_sample_records_reproducible(ctx: CodebookContext) -> None:
    rs = [cast_record(_make_record(ctx, f"W{i}")) for i in range(20)]
    a = [r.work_id for r in sample_records(rs, 7, seed=123)]
    b = [r.work_id for r in sample_records(rs, 7, seed=123)]
    c = [r.work_id for r in sample_records(rs, 7, seed=456)]
    assert a == b
    assert a != c


def test_short_work_id_openalex_url() -> None:
    assert short_work_id("https://openalex.org/W3213722062") == "W3213722062"


def test_unique_sheet_name_collision_suffix() -> None:
    used: set[str] = set()
    assert _unique_sheet_name("W1", used) == "W1"
    assert _unique_sheet_name("W1", used) == "W1_2"


# ── per-work tab field table ────────────────────────────────────────────────


def test_work_sheet_lists_every_codebook_field(
    tmp_path: Path, ctx: CodebookContext
) -> None:
    """The per-work tab lists each codebook field with value/context/reviewer
    columns, driven off the codebook field list (in spec order)."""
    work = _make_work("https://openalex.org/W1", title="The Title")
    record = _make_record(
        ctx, "https://openalex.org/W1", integration_scheme="RK4"
    )
    records_path, catalogue_path, output_path = _write_inputs(
        tmp_path, [work], [record]
    )

    export_review_xlsx(
        records_path, catalogue_path, output_path, ctx=ctx, n_subset=None
    )
    wb = load_workbook(output_path)
    ws = wb["W1"]

    # Field table header on row 12.
    assert ws["A12"].value == "field"
    assert ws["B12"].value == "value"
    assert ws["C12"].value == "context"
    assert ws["D12"].value == "reviewer_correction"

    # The field rows follow the codebook's spec field order.
    spec_field_names = [f.name for f in ctx.spec.fields]
    listed: list[str] = []
    row = 13
    while ws.cell(row=row, column=1).value in spec_field_names:
        listed.append(ws.cell(row=row, column=1).value)
        row += 1
    assert listed == spec_field_names
    # The first field carries its extracted value and a reviewer placeholder.
    first_value_row = 13 + spec_field_names.index("integration_scheme")
    assert ws.cell(row=first_value_row, column=2).value == "RK4"
    assert "<correct the value" in ws.cell(row=first_value_row, column=4).value


def test_work_sheet_follows_non_default_codebook(tmp_path: Path) -> None:
    """A non-default codebook's field set drives the per-work tab."""
    alt_ctx = CodebookContext.from_spec(_ALT_CODEBOOK)
    work = _make_work("https://openalex.org/W1")
    record = _make_record(alt_ctx, "https://openalex.org/W1", alpha="A!", beta="B!")
    records_path, catalogue_path, output_path = _write_inputs(
        tmp_path, [work], [record]
    )

    export_review_xlsx(
        records_path, catalogue_path, output_path, ctx=alt_ctx, n_subset=None
    )
    wb = load_workbook(output_path)
    ws = wb["W1"]

    # Only the alt codebook's two value fields appear, in order.
    assert ws["A13"].value == "alpha"
    assert ws["B13"].value == "A!"
    assert ws["A14"].value == "beta"
    assert ws["B14"].value == "B!"
    # No default-codebook field leaked in.
    assert ws["A15"].value == "reviewer_remarks"


# ── overview / Index sheet ──────────────────────────────────────────────────


def test_index_sheet_links_to_each_tab(tmp_path: Path, ctx: CodebookContext) -> None:
    works = [
        _make_work("https://openalex.org/W1"),
        _make_work("https://openalex.org/W2"),
    ]
    records = [
        _make_record(ctx, "https://openalex.org/W1"),
        _make_record(ctx, "https://openalex.org/W2"),
    ]
    records_path, catalogue_path, output_path = _write_inputs(tmp_path, works, records)

    export_review_xlsx(
        records_path, catalogue_path, output_path, ctx=ctx, n_subset=None
    )
    wb = load_workbook(output_path)
    assert wb.sheetnames == ["Index", "W1", "W2"]
    index = wb["Index"]
    assert index["A5"].value == "work_id"
    assert index["F5"].value == "sheet"
    link = index["F6"]
    assert link.value == "W1"
    assert link.hyperlink is not None
    assert link.hyperlink.location == "'W1'!A1"


# ── subset / seed ───────────────────────────────────────────────────────────


def test_subset_samples_reproducibly(tmp_path: Path, ctx: CodebookContext) -> None:
    works = [_make_work(f"https://openalex.org/W{i}") for i in range(10)]
    records = [_make_record(ctx, f"https://openalex.org/W{i}") for i in range(10)]
    records_path, catalogue_path, output_path = _write_inputs(tmp_path, works, records)

    count = export_review_xlsx(
        records_path, catalogue_path, output_path, ctx=ctx, n_subset=4, seed=7
    )
    assert count == 4
    wb = load_workbook(output_path)
    per_work = wb.sheetnames[1:]
    assert len(per_work) == 4
    indices = [int(name[1:]) for name in per_work]
    assert indices == sorted(indices)

    output_path2 = tmp_path / "review2.xlsx"
    export_review_xlsx(
        records_path, catalogue_path, output_path2, ctx=ctx, n_subset=4, seed=7
    )
    wb2 = load_workbook(output_path2)
    assert wb2.sheetnames[1:] == per_work


# ── missing catalogue ───────────────────────────────────────────────────────


def test_raises_on_missing_work(tmp_path: Path, ctx: CodebookContext) -> None:
    works = [_make_work("https://openalex.org/W1")]
    records = [_make_record(ctx, "https://openalex.org/W_missing")]
    records_path, catalogue_path, output_path = _write_inputs(tmp_path, works, records)

    with pytest.raises(ValueError, match="W_missing"):
        export_review_xlsx(
            records_path, catalogue_path, output_path, ctx=ctx, n_subset=None
        )


# ── CLI smoke ───────────────────────────────────────────────────────────────


def test_cli_default_output(tmp_path: Path, ctx: CodebookContext) -> None:
    from laglitsynth.cli import main

    works = [_make_work("https://openalex.org/W1")]
    records = [_make_record(ctx, "https://openalex.org/W1")]
    records_path, catalogue_path, _ = _write_inputs(tmp_path, works, records)

    main(
        [
            "extraction-codebook-export",
            "--records",
            str(records_path),
            "--catalogue",
            str(catalogue_path),
        ]
    )

    assert (tmp_path / "review.xlsx").exists()

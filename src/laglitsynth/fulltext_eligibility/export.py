"""XLSX export of stage 7 eligibility verdicts for human review.

Single subcommand ``fulltext-eligibility-export`` that writes a workbook with
an ``Index`` sheet plus one tab per sampled work. The Index sheet is a light
navigation aid: one row per sampled work with bibliographic columns, the
``eligible`` / ``llm_reason`` verdict inline, and a hyperlink to the per-work
tab. Per-work tabs are the working surface: a vertical ``Field | Value``
layout carrying the bibliographic block, the eligibility criterion pulled
verbatim from ``eligibility-meta.json`` so the reviewer sees the same question
the LLM saw, reviewer placeholders (``reviewer_eligible`` /
``reviewer_reason``), and a collapsed LLM-verdict block.

``--n-subset`` + ``--subset-seed`` draw a reproducible random sample;
``n_subset >= len(verdicts)`` (or unset) emits the full set in verdict order.
"""

from __future__ import annotations

import argparse
import json
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.worksheet import Worksheet

from laglitsynth.catalogue_fetch.models import Work
from laglitsynth.fulltext_eligibility.models import EligibilityMeta, EligibilityVerdict
from laglitsynth.io import read_jsonl

_SHEET_NAME_MAX_LEN = 31
_FORBIDDEN_SHEET_CHARS = set("/\\?*[]:")
_DEFAULT_SUBSET_SEED = 0

_REVIEWER_ELIGIBLE_PLACEHOLDER = (
    "<enter yes/no — does this paper meet the criterion above, based on the "
    "full text? — based on info provided above only>"
)
_REVIEWER_REASON_PLACEHOLDER = (
    "<give short (max 3 sentences) reason for the eligibility decision>"
)
_REVIEWER_REMARKS_PLACEHOLDER = (
    "<any general feedback on this record — paper quality, "
    "metadata gaps, oddities — leave blank if none>"
)

_HYPERLINK_FONT = Font(color="0563C1", underline="single")
_BOLD = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="E0E0E0")


def sample_verdicts(
    verdicts: list[EligibilityVerdict],
    n_subset: int | None,
    seed: int,
) -> list[EligibilityVerdict]:
    """Return a subset of verdicts, preserving their original order.

    ``n_subset is None`` or ``n_subset >= len(verdicts)`` returns the
    full list. Otherwise draw ``n_subset`` work_ids uniformly with
    ``random.Random(seed)`` and emit the matching verdicts in the
    order they appear in ``verdicts``.
    """
    if n_subset is None or n_subset >= len(verdicts):
        return list(verdicts)
    rng = random.Random(seed)
    chosen_ids = set(rng.sample([v.work_id for v in verdicts], k=n_subset))
    return [v for v in verdicts if v.work_id in chosen_ids]


def short_work_id(work_id: str) -> str:
    """Trailing OpenAlex id, sanitised for Excel sheet naming.

    Strips host/path prefix, replaces characters Excel forbids in
    sheet names with ``_``, truncates to 31 chars.
    """
    name = work_id.rsplit("/", 1)[-1] if "/" in work_id else work_id
    name = "".join("_" if c in _FORBIDDEN_SHEET_CHARS else c for c in name)
    name = name.strip("'")
    return name[:_SHEET_NAME_MAX_LEN] or "work"


def _unique_sheet_name(base: str, used: set[str]) -> str:
    if base not in used:
        used.add(base)
        return base
    for i in range(2, 10**6):
        suffix = f"_{i}"
        candidate = f"{base[: _SHEET_NAME_MAX_LEN - len(suffix)]}{suffix}"
        if candidate not in used:
            used.add(candidate)
            return candidate
    raise RuntimeError(f"Could not resolve unique sheet name for {base!r}")


def _authors_string(work: Work) -> str:
    names = [a.author.display_name for a in work.authorships]
    return ", ".join(names)


def _journal_name(work: Work) -> str:
    if work.primary_location and work.primary_location.source:
        return work.primary_location.source.display_name or ""
    return ""


def _openalex_external_url(work_id: str) -> str:
    """openalex IDs come back as full URLs already; just normalise."""
    if work_id.startswith("http://") or work_id.startswith("https://"):
        return work_id
    trailing = work_id.rsplit("/", 1)[-1]
    return f"https://openalex.org/{trailing}"


def _doi_url(doi: str) -> str:
    if doi.startswith("http://") or doi.startswith("https://"):
        return doi
    return f"https://doi.org/{doi}"


def _eligible_label(eligible: bool | None) -> str:
    """Render the tri-state ``eligible`` verdict for inline display."""
    if eligible is True:
        return "yes"
    if eligible is False:
        return "no"
    return ""


def build_index_sheet(
    ws: Worksheet, items: list[tuple[EligibilityVerdict, Work, str]]
) -> None:
    """Index sheet with reviewer-identity header rows above the table.

    The eligibility verdict (``eligible`` / ``llm_reason``) is shown
    inline as a navigation aid — the reviewer scans for tabs to dig
    into; the per-work tab is the working surface.
    """
    ws.title = "Index"

    ws.cell(row=1, column=1, value="reviewer_name").font = _BOLD
    ws.cell(row=1, column=2, value="<enter reviewer name here>")
    ws.cell(row=2, column=1, value="reviewer_email").font = _BOLD
    ws.cell(row=2, column=2, value="<enter reviewer email here>")
    ws.cell(row=3, column=1, value="review_date").font = _BOLD
    ws.cell(row=3, column=2, value="<YYYY-MM-DD>")

    headers = (
        "work_id",
        "title",
        "authors",
        "journal",
        "year",
        "eligible",
        "llm_reason",
        "sheet",
    )
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=col, value=header)
        cell.font = _BOLD
        cell.fill = _HEADER_FILL

    for row_offset, (verdict, work, sheet_name) in enumerate(items):
        row = 6 + row_offset
        ws.cell(row=row, column=1, value=verdict.work_id)
        ws.cell(row=row, column=2, value=work.title or "")
        ws.cell(row=row, column=3, value=_authors_string(work))
        ws.cell(row=row, column=4, value=_journal_name(work))
        ws.cell(row=row, column=5, value=work.publication_year)
        ws.cell(row=row, column=6, value=_eligible_label(verdict.eligible))
        ws.cell(row=row, column=7, value=verdict.reason or "")
        link_cell = ws.cell(row=row, column=8, value=sheet_name)
        link_cell.hyperlink = Hyperlink(
            ref=link_cell.coordinate,
            location=f"'{sheet_name}'!A1",
            display=sheet_name,
        )
        link_cell.font = _HYPERLINK_FONT

    # Freeze identity rows + table header.
    ws.freeze_panes = "A6"
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 36
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["E"].width = 8
    ws.column_dimensions["F"].width = 10
    ws.column_dimensions["G"].width = 50
    ws.column_dimensions["H"].width = 20


def build_work_sheet(
    ws: Worksheet,
    verdict: EligibilityVerdict,
    work: Work,
    *,
    criterion: str,
    llm_meta: dict[str, object],
) -> None:
    """Per-work sheet: bibliographic block, criterion, reviewer placeholders, LLM (collapsed).

    ``criterion`` is the eligibility-criteria system prompt rendered
    verbatim so the reviewer decides against the same question the LLM
    saw. ``llm_meta`` carries the model/temperature/prompt_sha256 from
    ``EligibilityMeta.llm`` so the reviewer can audit the LLM's run.
    """
    wrap_top_left = Alignment(wrap_text=True, vertical="top", horizontal="left")
    top_aligned = Alignment(vertical="top")

    # Row 1: back-to-Index hyperlink
    back_cell = ws.cell(row=1, column=1, value="← back to Index")
    back_cell.hyperlink = Hyperlink(
        ref=back_cell.coordinate,
        location="'Index'!A1",
        display="← back to Index",
    )
    back_cell.font = _HYPERLINK_FONT

    # Row 2: blank

    # Bibliographic block (rows 3-8)
    biblio_rows: list[tuple[str, object]] = [
        ("title", work.title or ""),
        ("authors", _authors_string(work)),
        ("journal", _journal_name(work)),
        ("publication_year", work.publication_year),
        ("doi", None),  # filled separately to attach a hyperlink
        ("openalex", None),  # ditto
    ]
    for offset, (field, value) in enumerate(biblio_rows):
        row = 3 + offset
        f_cell = ws.cell(row=row, column=1, value=field)
        f_cell.font = _BOLD
        f_cell.alignment = top_aligned
        if field == "doi":
            doi = work.doi or ""
            v_cell = ws.cell(row=row, column=2, value=doi)
            if doi:
                v_cell.hyperlink = Hyperlink(
                    ref=v_cell.coordinate,
                    target=_doi_url(doi),
                    display=doi,
                )
                v_cell.font = _HYPERLINK_FONT
            v_cell.alignment = wrap_top_left
        elif field == "openalex":
            url = _openalex_external_url(work.id)
            v_cell = ws.cell(row=row, column=2, value=url)
            v_cell.hyperlink = Hyperlink(
                ref=v_cell.coordinate,
                target=url,
                display=url,
            )
            v_cell.font = _HYPERLINK_FONT
            v_cell.alignment = wrap_top_left
        else:
            v_cell = ws.cell(row=row, column=2, value=value)
            v_cell.alignment = wrap_top_left

    # Row 9: blank

    # Criterion + reviewer decision (rows 10-13)
    crit_label = ws.cell(row=10, column=1, value="criterion")
    crit_label.font = _BOLD
    crit_label.alignment = top_aligned
    crit_value = ws.cell(row=10, column=2, value=criterion)
    crit_value.alignment = wrap_top_left

    re_label = ws.cell(row=11, column=1, value="reviewer_eligible")
    re_label.font = _BOLD
    re_label.alignment = top_aligned
    re_value = ws.cell(row=11, column=2, value=_REVIEWER_ELIGIBLE_PLACEHOLDER)
    re_value.alignment = wrap_top_left

    rr_label = ws.cell(row=12, column=1, value="reviewer_reason")
    rr_label.font = _BOLD
    rr_label.alignment = top_aligned
    rr_value = ws.cell(row=12, column=2, value=_REVIEWER_REASON_PLACEHOLDER)
    rr_value.alignment = wrap_top_left

    rmk_label = ws.cell(row=13, column=1, value="reviewer_remarks")
    rmk_label.font = _BOLD
    rmk_label.alignment = top_aligned
    rmk_value = ws.cell(row=13, column=2, value=_REVIEWER_REMARKS_PLACEHOLDER)
    rmk_value.alignment = wrap_top_left

    # Row 14: blank

    # LLM details — header on row 15 visible, rows 16-21 collapsed.
    llm_header = ws.cell(row=15, column=1, value="LLM details (expand to peek)")
    llm_header.font = _BOLD
    llm_header.alignment = top_aligned

    llm_eligible_label = ws.cell(row=16, column=1, value="llm_eligible")
    llm_eligible_label.font = _BOLD
    llm_eligible_label.alignment = top_aligned
    ws.cell(row=16, column=2, value=_eligible_label(verdict.eligible))

    llm_reason_label = ws.cell(row=17, column=1, value="llm_reason")
    llm_reason_label.font = _BOLD
    llm_reason_label.alignment = top_aligned
    llm_reason_value = ws.cell(row=17, column=2, value=verdict.reason or "")
    llm_reason_value.alignment = wrap_top_left

    llm_model_label = ws.cell(row=18, column=1, value="llm_model")
    llm_model_label.font = _BOLD
    llm_model_label.alignment = top_aligned
    ws.cell(row=18, column=2, value=str(llm_meta.get("model", "")))

    llm_temp_label = ws.cell(row=19, column=1, value="llm_temperature")
    llm_temp_label.font = _BOLD
    llm_temp_label.alignment = top_aligned
    ws.cell(row=19, column=2, value=llm_meta.get("temperature"))

    llm_sha_label = ws.cell(row=20, column=1, value="llm_prompt_sha256")
    llm_sha_label.font = _BOLD
    llm_sha_label.alignment = top_aligned
    ws.cell(row=20, column=2, value=str(llm_meta.get("prompt_sha256", "")))

    llm_raw_label = ws.cell(row=21, column=1, value="llm_raw_response")
    llm_raw_label.font = _BOLD
    llm_raw_label.alignment = top_aligned
    llm_raw_value = ws.cell(row=21, column=2, value=verdict.raw_response or "")
    llm_raw_value.alignment = wrap_top_left

    # Collapse the LLM rows; +/- button next to the visible header row.
    ws.sheet_properties.outlinePr.summaryBelow = False
    for row_idx in range(16, 22):
        ws.row_dimensions[row_idx].outline_level = 1
        ws.row_dimensions[row_idx].hidden = True

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110
    ws.freeze_panes = "B1"


def _load_meta(meta_path: Path | None) -> tuple[str, dict[str, object]]:
    """Return ``(criterion, llm_meta_dict)`` from an eligibility-meta.json file.

    Falls back to placeholders when ``meta_path`` is None or missing —
    the export still works, just without the criterion / LLM
    fingerprint. ``eligibility-meta.json`` may pre-date the addition of
    ``criterion`` to ``EligibilityMeta``, in which case the criterion
    shows a stub.
    """
    if meta_path is None or not meta_path.exists():
        print(
            f"WARNING: eligibility-meta.json not found at {meta_path}; "
            f"per-work sheets will show '<eligibility criterion not available>'. "
            f"Pass --meta to point at the right file.",
            file=sys.stderr,
        )
        return ("<eligibility criterion not available>", {})

    raw = json.loads(meta_path.read_text())
    # Use EligibilityMeta to validate, but tolerate older meta files.
    try:
        meta = EligibilityMeta.model_validate(raw)
        criterion = meta.criterion or "<eligibility criterion not recorded in meta>"
        llm_meta: dict[str, object] = {
            "model": meta.llm.model,
            "temperature": meta.llm.temperature,
            "prompt_sha256": meta.llm.prompt_sha256,
        }
        return (criterion, llm_meta)
    except Exception:
        criterion = raw.get("criterion") or "<eligibility criterion not recorded in meta>"
        llm_meta = raw.get("llm", {}) or {}
        return (criterion, llm_meta)


def export_review_xlsx(
    verdicts_path: Path,
    catalogue_path: Path,
    output_path: Path,
    *,
    n_subset: int | None,
    seed: int = _DEFAULT_SUBSET_SEED,
    meta_path: Path | None = None,
) -> int:
    """Build and save the review workbook, returning the per-work sheet count.

    Raises ``ValueError`` naming the first ``work_id`` present in the
    verdicts file but absent from the catalogue.
    """
    catalogue: dict[str, Work] = {w.id: w for w in read_jsonl(catalogue_path, Work)}
    verdicts = list(read_jsonl(verdicts_path, EligibilityVerdict))
    for verdict in verdicts:
        if verdict.work_id not in catalogue:
            raise ValueError(
                f"work_id {verdict.work_id!r} in {verdicts_path} not found in "
                f"catalogue {catalogue_path}"
            )

    selected = sample_verdicts(verdicts, n_subset, seed)
    criterion, llm_meta = _load_meta(meta_path)

    wb = Workbook()
    index_ws = wb.active
    if index_ws is None:
        raise RuntimeError("openpyxl Workbook() did not produce a default worksheet")

    used_names: set[str] = set()
    items: list[tuple[EligibilityVerdict, Work, str]] = []
    for verdict in selected:
        sheet_name = _unique_sheet_name(short_work_id(verdict.work_id), used_names)
        ws = wb.create_sheet(title=sheet_name)
        build_work_sheet(
            ws,
            verdict,
            catalogue[verdict.work_id],
            criterion=criterion,
            llm_meta=llm_meta,
        )
        items.append((verdict, catalogue[verdict.work_id], sheet_name))

    build_index_sheet(index_ws, items)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return len(selected)


def build_subparser(
    subparsers: argparse._SubParsersAction[argparse.ArgumentParser],
) -> argparse.ArgumentParser:
    parser = subparsers.add_parser(
        "fulltext-eligibility-export",
        help="Export stage 7 eligibility verdicts + catalogue to an XLSX review workbook.",
    )
    parser.add_argument(
        "--verdicts",
        type=Path,
        required=True,
        help="Path to verdicts.jsonl from fulltext-eligibility",
    )
    parser.add_argument(
        "--catalogue",
        type=Path,
        required=True,
        help="Path to the deduplicated catalogue JSONL",
    )
    parser.add_argument(
        "--meta",
        type=Path,
        default=None,
        help=(
            "Path to eligibility-meta.json. Default: <verdicts parent>/eligibility-meta.json. "
            "Embeds the eligibility criterion and LLM fingerprint into each per-work sheet."
        ),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help="Output file path. Default: <verdicts parent>/review.xlsx.",
    )
    parser.add_argument(
        "--n-subset",
        type=int,
        default=None,
        help="Random sample size; emit all when unset or >= total.",
    )
    parser.add_argument(
        "--subset-seed",
        type=int,
        default=_DEFAULT_SUBSET_SEED,
        help=f"Random seed for --n-subset (default: {_DEFAULT_SUBSET_SEED}).",
    )
    parser.set_defaults(run=run)
    return parser


def run(args: argparse.Namespace) -> None:
    verdicts_path: Path = args.verdicts

    output_path: Path = (
        args.output
        if args.output is not None
        else verdicts_path.parent / "review.xlsx"
    )
    meta_path: Path = (
        args.meta
        if args.meta is not None
        else verdicts_path.parent / "eligibility-meta.json"
    )
    count = export_review_xlsx(
        verdicts_path,
        args.catalogue,
        output_path,
        n_subset=args.n_subset,
        seed=args.subset_seed,
        meta_path=meta_path,
    )
    print(
        f"Wrote workbook with {count} per-work sheets to {output_path}",
        file=sys.stderr,
    )
